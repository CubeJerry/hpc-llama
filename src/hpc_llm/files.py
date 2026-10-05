"""Explicit attachment parsing, immutable provenance, and deterministic local tables.

No parser opens model-selected paths. The parent reads approved regular files,
then a killable child receives bytes. Browsing never reads file contents.
"""
from __future__ import annotations

import base64
import copy
import csv
import hashlib
import io
import json
import math
import mimetypes
import os
from pathlib import Path
import re
import stat
import subprocess
import sys
import tempfile
import threading
import zipfile
from typing import Any

from .contracts import AppError, Attachment, BackendCapabilities, SourceRecord, OutputArtifact, MAX_DOCUMENT_BYTES

MAX_INPUT = 16 * 1024 * 1024
MAX_TEXT = 2 * 1024 * 1024
MAX_ARCHIVE = 64 * 1024 * 1024
MAX_PAGES = 200
MAX_ROWS = 20000
MAX_COLS = 256
MAX_LISTING = 300
PARSER_SECONDS = 15
TEXT_EXTENSIONS = {'.txt', '.md', '.markdown', '.log', '.py', '.r', '.R', '.js', '.ts', '.tsx', '.jsx', '.c', '.h', '.cpp', '.hpp', '.sh', '.bash', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.conf', '.tex', '.sql', '.rst', '.fasta', '.fa', '.pdb', '.cif'}
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp'}


def clean_text(value: Any) -> str:
    """Remove terminal controls; Rich markup is escaped by the UI, not mutated here."""
    value = str(value)
    value = re.sub(r'\x1b\][^\x07]*(?:\x07|\x1b\\)', '', value)
    value = re.sub(r'\x1b\[[0-?]*[ -/]*[@-~]', '', value)
    return ''.join(c for c in value if c in '\n\t' or (ord(c) >= 32 and ord(c) != 127))


def clean_name(value: Any) -> str:
    return clean_text(value).replace('\n', ' ').replace('\t', ' ')


def _private_dir(path: Path) -> None:
    if path.is_symlink():
        raise AppError('permission', 'Private snapshot directory cannot be a symlink.')
    path.mkdir(mode=0o700, parents=True, exist_ok=True)
    st = path.stat()
    if st.st_uid != os.getuid() or not stat.S_ISDIR(st.st_mode):
        raise AppError('permission', 'Private snapshot directory must be owned by you.')
    if st.st_mode & 0o077:
        os.chmod(path, 0o700)


def _atomic_private(path: Path, data: bytes, overwrite: bool = True) -> None:
    if path.is_symlink():
        raise AppError('permission', 'Refusing to write through a symbolic link.')
    fd, temporary = tempfile.mkstemp(prefix='.writing-', dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        if overwrite:
            os.replace(temporary, path)
        else:
            os.link(temporary, path)
            os.unlink(temporary)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)



def _publish_bytes(parent_fd: int, name: str, data: bytes, overwrite: bool) -> None:
    """Publish private bytes atomically within an already verified directory."""
    temporary = f'.export-{os.getpid()}-{os.urandom(6).hex()}'
    fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600, dir_fd=parent_fd)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        if overwrite:
            os.rename(temporary, name, src_dir_fd=parent_fd, dst_dir_fd=parent_fd)
        else:
            # Linking, unlike a check then rename, cannot clobber a racing writer.
            os.link(temporary, name, src_dir_fd=parent_fd, dst_dir_fd=parent_fd, follow_symlinks=False)
    except FileExistsError:
        raise AppError('validation', 'Destination exists. Confirm overwrite explicitly or choose another name.') from None
    finally:
        try:
            os.unlink(temporary, dir_fd=parent_fd)
        except FileNotFoundError:
            pass


def _safe_root(path: Path) -> Path:
    path = Path(os.path.abspath(path.expanduser()))
    for ancestor in [*reversed(path.parents), path]:
        if ancestor.is_symlink():
            raise AppError('permission', 'Workspace root cannot contain symbolic links; choose its real path explicitly.')
    if not path.is_dir():
        raise AppError('validation', 'Workspace must be an existing directory.')
    return path


def _relative(root: Path, path: str | Path | None) -> Path:
    candidate = Path(path or root).expanduser()
    if not candidate.is_absolute():
        candidate = root / candidate
    candidate = Path(os.path.abspath(candidate))
    try:
        relative = candidate.relative_to(root)
    except ValueError:
        raise AppError('permission', 'Path is outside the selected workspace. Choose that workspace explicitly first.') from None
    if any(part.startswith('.') for part in relative.parts):
        raise AppError('permission', 'Hidden paths are excluded from attachments and browsing.')
    return relative


def _open_at(root: Path, relative: Path, *, directory: bool = False) -> int:
    """Walk using directory FDs and O_NOFOLLOW to close symlink check/open races."""
    flags = os.O_RDONLY | getattr(os, 'O_NOFOLLOW', 0) | getattr(os, 'O_NONBLOCK', 0)
    root_fd = os.open(root, flags | os.O_DIRECTORY)
    fd = root_fd
    try:
        parts = relative.parts
        for index, part in enumerate(parts):
            last = index == len(parts) - 1
            next_fd = os.open(part, flags | (os.O_DIRECTORY if not last or directory else 0), dir_fd=fd)
            os.close(fd)
            fd = next_fd
        result = fd
        fd = -1
        return result
    except OSError as exc:
        raise AppError('permission', 'Cannot open path safely; check permissions and avoid symbolic links.') from exc
    finally:
        if fd >= 0:
            os.close(fd)


def _read_approved(root: Path, path: str) -> tuple[bytes, Path]:
    relative = _relative(root, path)
    fd = _open_at(root, relative)
    try:
        st = os.fstat(fd)
        if not stat.S_ISREG(st.st_mode):
            raise AppError('permission', 'Only regular files may be attached; devices, FIFOs and sockets are refused.')
        if st.st_size > MAX_INPUT:
            raise AppError('parser_limit', 'File exceeds the 16 MiB input limit. Create a smaller explicit excerpt first.')
        with os.fdopen(fd, 'rb', closefd=False) as stream:
            data = stream.read(MAX_INPUT + 1)
        if len(data) > MAX_INPUT:
            raise AppError('parser_limit', 'File grew beyond the 16 MiB input limit during attachment.')
        return data, root / relative
    finally:
        os.close(fd)


def _run_parser(data: bytes, extension: str, selection: dict, vision: bool = False) -> dict:
    if len(data) > MAX_INPUT:
        raise AppError('parser_limit', 'Parser input exceeds 16 MiB.')
    payload = json.dumps({'data': base64.b64encode(data).decode(), 'extension': extension.lower(), 'selection': selection, 'vision': vision}).encode()
    try:
        result = subprocess.run([sys.executable, '-m', 'hpc_llm.files', '--parse'], input=payload,
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=PARSER_SECONDS,
                                env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, check=False)
    except subprocess.TimeoutExpired:
        raise AppError('parser_limit', 'Document parsing exceeded the 15-second limit. Select a smaller document or range.') from None
    if result.returncode != 0 or len(result.stdout) > 32 * 1024 * 1024:
        raise AppError('parser_limit', 'Document parser stopped at its memory, time or output limit.')
    try:
        parsed = json.loads(result.stdout)
    except (ValueError, UnicodeError):
        raise AppError('parser_limit', 'Document parser returned an invalid response.') from None
    if 'error' in parsed:
        raise AppError(parsed.get('code', 'validation'), parsed['error'])
    return parsed


def parse_public_pdf(data: bytes) -> dict:
    """The same bounded worker used for local PDFs; no URL/network operations."""
    return _run_parser(data, '.pdf', {})


def parse_public_html(data: bytes) -> dict:
    """Bounded local HTML extraction without JavaScript or embedded requests."""
    return _run_parser(data, '.html', {})


def _decode(data: bytes) -> str:
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        if data[:2] in {b'\xff\xfe', b'\xfe\xff'}:
            try:
                return data.decode('utf-16')
            except UnicodeDecodeError:
                pass
        raise AppError('validation', 'Text is not valid UTF-8 or BOM-marked UTF-16. Save a UTF-8 copy and attach it.') from None


def _range_indices(value: Any, total: int, label: str) -> list[int]:
    if value is None:
        return list(range(1, total + 1))
    if isinstance(value, list):
        numbers = value
    elif isinstance(value, str):
        numbers = []
        for piece in value.split(','):
            match = re.fullmatch(r'\s*(\d+)(?:\s*-\s*(\d+))?\s*', piece)
            if not match:
                raise AppError('validation', f'{label} must use numbers and ranges such as 1-3,5.')
            start, end = int(match[1]), int(match[2] or match[1])
            if end < start or end - start > MAX_ROWS:
                raise AppError('validation', f'Invalid {label} range.')
            numbers.extend(range(start, end + 1))
    else:
        raise AppError('validation', f'{label} must be a range string or list of integers.')
    if not numbers or any(isinstance(n, bool) or not isinstance(n, int) or n < 1 or n > total for n in numbers):
        raise AppError('validation', f'{label} falls outside 1-{total}.')
    return sorted(set(numbers))


def _compact_numbers(numbers: list[int]) -> str:
    ranges=[]
    for number in numbers:
        if ranges and number==ranges[-1][1]+1:
            ranges[-1][1]=number
        else:
            ranges.append([number,number])
    return ','.join(str(a) if a==b else f'{a}-{b}' for a,b in ranges)


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_ARCHIVE:
                raise AppError('parser_limit', 'Document archive exceeds 64 MiB decompressed or 10,000-entry limit.')
            for item in infos:
                if item.flag_bits & 1:
                    raise AppError('unsupported', 'Encrypted document archives are unsupported.')
                if item.file_size > 1024 * 1024 and item.file_size / max(1, item.compress_size) > 200:
                    raise AppError('parser_limit', 'Document has an unsafe decompression ratio.')
                if item.filename.lower().endswith(('vbaproject.bin', '.exe')):
                    raise AppError('unsupported', 'Macro-enabled documents are unsupported. Export a plain DOCX/XLSX copy.')
    except zipfile.BadZipFile:
        raise AppError('validation', 'Malformed office document archive.') from None


def _parse(data: bytes, extension: str, selection: dict, vision: bool) -> dict:
    if not isinstance(selection, dict):
        raise AppError('validation', 'Selection must be an object.')
    allowed = ({'lines','start_line','end_line'} if extension in TEXT_EXTENSIONS or extension == ''
               else {'start_row','end_row','columns'} if extension in {'.csv','.tsv'}
               else {'pointer'} if extension == '.json'
               else {'pages'} if extension == '.pdf'
               else {'sheet','range'} if extension == '.xlsx' else set())
    if set(selection) - allowed:
        raise AppError('validation', 'Unsupported selection fields for this file type: ' + ', '.join(sorted(set(selection)-allowed)))
    result: dict[str, Any] = {'sources': [], 'warnings': [], 'metadata': {}, 'status': 'ready'}
    used = 0

    def add(text: Any, locator: str, partial: bool = False) -> None:
        nonlocal used
        text = clean_text(text)
        available = MAX_TEXT - used
        encoded = text.encode('utf-8')
        if len(encoded) > available:
            text = encoded[:max(0, available)].decode('utf-8', errors='ignore')
            partial = True
            result['status'] = 'partial'
            if 'Extraction reached 2 MiB text limit; coverage is incomplete.' not in result['warnings']:
                result['warnings'].append('Extraction reached 2 MiB text limit; coverage is incomplete.')
        used += len(text.encode('utf-8'))
        if text.strip():
            result['sources'].append({'text': text, 'locator': locator, 'partial': partial})

    def add_table(name: str, headers: list, rows: list, refs: list, **extra: Any) -> None:
        headers = [clean_text(h) if h not in (None, '') else f'column_{i + 1}' for i, h in enumerate(headers)]
        seen: dict[str, int] = {}
        unique = []
        for header in headers:
            seen[header] = seen.get(header, 0) + 1
            unique.append(header if seen[header] == 1 else f'{header}_{seen[header]}')
        table = {'name': name, 'headers': unique, 'rows': rows, 'row_references': refs, 'coverage': 'all selected rows', **extra}
        result['metadata'].setdefault('tables', []).append(table)
        for offset in range(0, len(rows), 40):
            out = io.StringIO()
            writer = csv.writer(out, lineterminator='\n')
            writer.writerow(['source_reference', *unique])
            for ref, row in zip(refs[offset:offset + 40], rows[offset:offset + 40]):
                writer.writerow([ref, *row])
            add(out.getvalue(), f'{name}: {refs[offset]} to {refs[min(offset + 39, len(refs)-1)]}')
        if not rows:
            add(','.join(unique), f'{name}: headers (no selected data rows)')

    if extension in TEXT_EXTENSIONS or extension == '':
        text = _decode(data)
        if '\x00' in text:
            raise AppError('unsupported', 'Binary content is not supported as a text attachment.')
        lines = text.splitlines()
        numbers = _range_indices(selection.get('lines'), len(lines), 'Lines') if lines else []
        if 'start_line' in selection or 'end_line' in selection:
            numbers = _range_indices(f"{selection.get('start_line', 1)}-{selection.get('end_line',len(lines))}", len(lines), 'Lines')
        for offset in range(0, len(numbers), 80):
            chunk = numbers[offset:offset+80]
            # Non-contiguous selections keep explicit line numbers, rather than implying omitted lines were read.
            add('\n'.join(f'{n}: {lines[n-1]}' for n in chunk), 'lines ' + _compact_numbers(chunk))
        result['metadata'].update(total_lines=len(lines), selected_lines=len(numbers))
    elif extension in {'.csv', '.tsv'}:
        reader = csv.reader(io.StringIO(_decode(data)), delimiter='\t' if extension == '.tsv' else ',')
        try:
            headers = next(reader)
        except StopIteration:
            raise AppError('validation', 'Table is empty.') from None
        if len(headers) > MAX_COLS:
            raise AppError('parser_limit', 'Table exceeds 256 columns.')
        rows, refs = [], []
        start = int(selection.get('start_row', 2)); end = int(selection.get('end_row', MAX_ROWS + 1))
        if start < 2 or end < start or end - start >= MAX_ROWS:
            raise AppError('validation', 'CSV selection must be data rows 2 onward, at most 20,000 rows.')
        columns = selection.get('columns', headers)
        if not isinstance(columns, list) or any(c not in headers for c in columns):
            raise AppError('validation', 'Columns must be existing header names.')
        positions = [headers.index(c) for c in columns]
        total = 1
        for number, row in enumerate(reader, 2):
            total = number
            if len(row) > MAX_COLS:
                raise AppError('parser_limit', 'Table exceeds 256 columns.')
            if start <= number <= end:
                if len(row) != len(headers):
                    result['warnings'].append(f'Row {number} width differs from header; absent cells are missing.')
                rows.append([clean_text(row[c]) if c < len(row) else None for c in positions]); refs.append(f'row {number}')
            if number > MAX_ROWS + 1 and 'end_row' not in selection:
                result['status'] = 'partial'; result['warnings'].append('Table exceeds 20,000 rows; only selected first rows were ingested.'); break
            if number > end:
                break
        add_table('table', columns, rows, refs)
        result['metadata']['rows_scanned'] = total
    elif extension == '.json':
        try:
            document = json.loads(_decode(data),object_pairs_hook=_json_object,parse_constant=_invalid_json_constant)
        except (ValueError, RecursionError):
            raise AppError('validation', 'Malformed or overly deep JSON document.') from None
        pointer = selection.get('pointer', '')
        if not isinstance(pointer, str) or (pointer and not pointer.startswith('/')):
            raise AppError('validation', 'JSON pointer must be empty or start with /.')
        selected = document
        try:
            for part in pointer.split('/')[1:]:
                part = part.replace('~1','/').replace('~0','~')
                selected = selected[int(part)] if isinstance(selected, list) else selected[part]
        except (KeyError, IndexError, TypeError, ValueError):
            raise AppError('validation', 'Selected JSON pointer does not exist.') from None
        pretty = json.dumps(selected, indent=2, ensure_ascii=False)
        for offset in range(0, len(pretty), 6000):
            add(pretty[offset:offset+6000], f'JSON pointer {pointer or "/ (root)"}; characters {offset+1}-{min(offset+6000,len(pretty))}')
        result['metadata'].update(json_pointer=pointer, json_type=type(selected).__name__)
        if isinstance(selected, list) and selected and all(isinstance(r, dict) for r in selected[:MAX_ROWS]):
            headers = list(dict.fromkeys(k for row in selected[:MAX_ROWS] for k in row))
            if len(headers) <= MAX_COLS:
                rows = [[json.dumps(row[h], ensure_ascii=False) if isinstance(row.get(h), (dict,list)) else row.get(h) for h in headers] for row in selected[:MAX_ROWS]]
                result['metadata']['tables'] = [{'name': 'json', 'headers': headers, 'rows': rows, 'row_references': [f'{pointer}/{i}' for i in range(len(rows))], 'coverage': f'first {len(rows)} selected array elements' if len(selected)>MAX_ROWS else 'all selected array elements'}]
            if len(selected) > MAX_ROWS:
                result['status'] = 'partial'; result['warnings'].append('Table operations cover only the first 20,000 selected JSON elements.')
    elif extension == '.pdf':
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise AppError('unsupported', 'Encrypted PDFs are unsupported. Attach a decrypted copy you are allowed to read.')
            total = len(reader.pages)
            numbers = _range_indices(selection.get('pages'), total, 'Pages')
            if len(numbers) > MAX_PAGES:
                raise AppError('parser_limit', 'Select at most 200 PDF pages.')
            missing = []
            for number in numbers:
                try:
                    text = reader.pages[number-1].extract_text() or ''
                except Exception:
                    text = ''
                if not text.strip():
                    missing.append(number)
                else:
                    for offset in range(0, len(text), 6000):
                        add(text[offset:offset+6000], f'page {number}' + (f'; characters {offset+1}-{min(offset+6000,len(text))}' if len(text)>6000 else ''))
            result['metadata'].update(pages=total, selected_pages=numbers, missing_pages=missing)
            result['warnings'].append('PDF extraction covers selectable text; diagrams and embedded table layout are not reliably read. No OCR was performed.')
            if missing:
                result['status'] = 'partial'; result['warnings'].append('No extractable text on pages ' + ', '.join(map(str, missing)) + ' (scanned, image-only or unextractable).')
        except AppError:
            raise
        except Exception:
            raise AppError('validation', 'Malformed or unreadable PDF document.') from None
    elif extension == '.docx':
        _check_archive(data)
        from docx import Document
        document = Document(io.BytesIO(data))
        paragraphs = document.paragraphs
        for index, paragraph in enumerate(paragraphs, 1):
            if paragraph.text.strip():
                add(paragraph.text, f'paragraph {index}' + (f' ({paragraph.style.name})' if paragraph.style else ''))
        for table_index, table in enumerate(document.tables, 1):
            if len(table.rows) > MAX_ROWS or len(table.columns) > MAX_COLS:
                raise AppError('parser_limit', 'DOCX table exceeds row/column limits.')
            rows = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
            if rows:
                add_table(f'table {table_index}', rows[0], rows[1:], [f'table {table_index} row {n}' for n in range(2,len(rows)+1)], cell_provenance='row and column index (1-based); first row used as headers')
        result['metadata']['paragraphs'] = len(paragraphs)
        result['warnings'].append('DOCX locations are paragraph/table/cell positions, not fixed page numbers. External relationships are not fetched.')
    elif extension == '.xlsx':
        _check_archive(data)
        import openpyxl
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            names = [selection['sheet']] if selection.get('sheet') else formulas.sheetnames
            if len(names) > 100:
                raise AppError('parser_limit', 'Select one sheet; workbook has more than 100 sheets.')
            missing_cache = []
            result['metadata']['sheets'] = formulas.sheetnames
            for name in names:
                if name not in formulas.sheetnames:
                    raise AppError('validation', 'Selected spreadsheet sheet does not exist.')
                sheet = formulas[name]; cached = values[name]
                if selection.get('range'):
                    try:
                        mincol, minrow, maxcol, maxrow = range_boundaries(selection['range'])
                    except (ValueError, TypeError):
                        raise AppError('validation', 'Spreadsheet range must be A1 notation, such as A1:D20.') from None
                    if None in (mincol,minrow,maxcol,maxrow):
                        raise AppError('validation', 'Use a bounded cell range, such as A1:D20.')
                else:
                    mincol, minrow, maxcol, maxrow = 1, 1, sheet.max_column or 1, sheet.max_row or 1
                if maxrow < minrow or maxcol < mincol or maxrow-minrow >= MAX_ROWS or maxcol-mincol >= MAX_COLS:
                    raise AppError('parser_limit', 'Select no more than 20,000 rows and 256 columns per sheet.')
                rows = []; refs = []; formula_records = []
                f_rows = sheet.iter_rows(min_row=minrow,max_row=maxrow,min_col=mincol,max_col=maxcol)
                v_rows = cached.iter_rows(min_row=minrow,max_row=maxrow,min_col=mincol,max_col=maxcol)
                for n,(fr,vr) in enumerate(zip(f_rows,v_rows),minrow):
                    row = []
                    for c,v in zip(fr,vr):
                        value = v.value
                        if c.data_type == 'f':
                            formula_records.append({'cell': c.coordinate, 'formula': c.value, 'cached_value': _scalar(value)})
                            if value is None:
                                missing_cache.append(f'{name}!{c.coordinate}')
                        row.append(_scalar(value))
                    rows.append(row); refs.append(f'{name}!{get_column_letter(mincol)}{n}:{get_column_letter(maxcol)}{n}')
                add_table(name, rows[0] if rows else [], rows[1:], refs[1:], selected_range=f'{get_column_letter(mincol)}{minrow}:{get_column_letter(maxcol)}{maxrow}', formulas=formula_records, header_reference=refs[0] if refs else '')
                if formula_records:
                    add(json.dumps(formula_records, ensure_ascii=False), f'{name}: formulas and cached values')
            result['metadata']['missing_formula_caches'] = missing_cache
            result['warnings'].append('XLSX formulas are not recalculated; cached results may be stale. External links and macros are not executed. Only worksheet cell data is extracted; embedded charts, images and drawing layout are not interpreted as vision.')
            if missing_cache:
                result['status'] = 'partial'; result['warnings'].append(f'{len(missing_cache)} formula cells have no cached result; treated as missing, never zero.')
        finally:
            formulas.close(); values.close()
    elif extension in {'.html', '.htm'}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_decode(data), 'html.parser')
        for tag in soup(['script','style','noscript','iframe','object','embed','template']):
            tag.decompose()
        text = soup.get_text('\n', strip=True)
        lines = text.splitlines()
        for offset in range(0,len(lines),80):
            add('\n'.join(lines[offset:offset+80]), f'extracted text lines {offset+1}-{min(offset+80,len(lines))}')
        result['metadata']['title'] = soup.title.get_text(' ', strip=True) if soup.title else ''
        result['warnings'].append('HTML parsed locally without JavaScript or resource loading; references are extracted-text lines.')
    elif extension in IMAGE_EXTENSIONS:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 20_000_000
        try:
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                if width * height > 20_000_000 or width > 20000 or height > 20000:
                    raise AppError('parser_limit', 'Image exceeds 20 megapixel or 20,000-pixel dimension limit.')
                image.verify()
        except AppError:
            raise
        except Exception:
            raise AppError('validation', 'Malformed, unsupported or oversized image.') from None
        result['metadata'].update(width=width,height=height)
        if not vision:
            result['status'] = 'unsupported'; result['warnings'].append('Image requires a model with verified vision support and a matching projector. No image was sent to the model.')
        elif len(data) > 5 * 1024 * 1024:
            raise AppError('parser_limit', 'Vision attachments are limited to 5 MiB; resize locally first.')
        else:
            mime = {'.jpg':'image/jpeg','.jpeg':'image/jpeg','.png':'image/png','.webp':'image/webp'}[extension]
            result['metadata']['image_data_uri'] = 'data:' + mime + ';base64,' + base64.b64encode(data).decode()
            add(f'Attached image ({width} × {height} pixels); visual content is supplied separately to the capable model.', 'image')
    elif extension == '.xls':
        raise AppError('unsupported', 'Legacy XLS workbooks are unsupported. Save a separate XLSX copy in Excel or LibreOffice, then attach it.')
    else:
        raise AppError('unsupported', 'Unsupported file type. Attach text, CSV/TSV, JSON, PDF, DOCX, XLSX, HTML, PNG, JPEG or WebP.')
    if not result['sources'] and result['status'] == 'ready':
        result['warnings'].append('No text content was extracted from this selection.')
    return result


def _json_object(pairs):
    result={}
    for key,value in pairs:
        if key in result:
            raise AppError('validation','JSON contains duplicate object keys; create an unambiguous copy.')
        result[key]=value
    return result


def _invalid_json_constant(value):
    raise AppError('validation','JSON contains non-finite numbers that are not valid JSON.')


def _scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str,int,bool)):
        return clean_text(value) if isinstance(value,str) else value
    if isinstance(value,float):
        return value if math.isfinite(value) else None
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return clean_text(value)


class FileService:
    def __init__(self, workspace: Path, snapshot_dir: Path, capabilities=None):
        self.workspace = _safe_root(workspace)
        self.snapshot_dir = Path(snapshot_dir)
        _private_dir(self.snapshot_dir)
        self.capabilities = capabilities or BackendCapabilities()
        self.attachments: dict[str, Attachment] = {}
        self._source_counter = 0
        self._lock = threading.RLock()

    def browse(self, path=None, query='') -> list[dict]:
        relative = _relative(self.workspace, path)
        fd = _open_at(self.workspace, relative, directory=True)
        entries = []
        try:
            # Bounded scan, not recursive search: at most 5,000 names inspected and 300 displayed.
            with os.scandir(fd) as iterator:
                for index, entry in enumerate(iterator):
                    if index >= 5000:
                        entries.append({'name':'More entries omitted; narrow this directory or filter.', 'kind':'notice', 'truncated':True}); break
                    if entry.name.startswith('.') or query.casefold() not in entry.name.casefold():
                        continue
                    info = entry.stat(follow_symlinks=False)
                    kind = 'directory' if stat.S_ISDIR(info.st_mode) else 'file' if stat.S_ISREG(info.st_mode) else 'symlink' if stat.S_ISLNK(info.st_mode) else 'special'
                    entries.append({'name':clean_name(entry.name), 'path':str(self.workspace / relative / entry.name), 'kind':kind, 'size_bytes':info.st_size, 'selectable':kind in {'file','directory'}})
                    if len(entries) >= MAX_LISTING:
                        entries.append({'name':'Listing limited to 300 entries; filter by name.', 'kind':'notice', 'truncated':True}); break
        except OSError:
            raise AppError('permission', 'Cannot list this directory.') from None
        finally:
            os.close(fd)
        return sorted(entries,key=lambda x:(x['kind']!='directory',x['name'].casefold()))

    def attach(self, path: str, selection: dict | None = None) -> Attachment:
        selection = selection or {}
        if len(json.dumps(selection)) > 8192:
            raise AppError('validation', 'Selection is too large.')
        data, approved_path = _read_approved(self.workspace, path)
        fingerprint = hashlib.sha256(data).hexdigest()
        identity = hashlib.sha256((fingerprint + str(approved_path) + json.dumps(selection,sort_keys=True)).encode()).hexdigest()[:20]
        attachment_id = 'A' + identity
        with self._lock:
            if attachment_id in self.attachments:
                return self.attachments[attachment_id].model_copy(deep=True)
            vision = bool(self.capabilities.get('vision',False) if isinstance(self.capabilities,dict) else self.capabilities.vision)
            parsed = _run_parser(data, approved_path.suffix, selection, vision)
            sources = []
            for source in parsed['sources']:
                self._source_counter += 1
                sources.append(SourceRecord(id=f'F{self._source_counter}',title=clean_name(approved_path.name),kind='file',attachment_id=attachment_id,fingerprint=fingerprint,**source))
            snapshot_path = self.snapshot_dir / (attachment_id + '-' + os.urandom(6).hex() + '.json')
            attachment = Attachment(id=attachment_id,name=clean_name(approved_path.name),path=str(approved_path),media_type=mimetypes.guess_type(approved_path.name)[0] or 'application/octet-stream',size_bytes=len(data),fingerprint=fingerprint,selection=selection,status=parsed['status'],warnings=parsed['warnings'],sources=sources,snapshot_path=str(snapshot_path),metadata=parsed['metadata'])
            _atomic_private(snapshot_path,attachment.model_dump_json().encode(),overwrite=False)
            self.attachments[attachment_id] = attachment
            return attachment.model_copy(deep=True)

    def restore(self, attachments) -> None:
        with self._lock:
            values = attachments.values() if isinstance(attachments,dict) else attachments
            for item in values:
                attachment = item if isinstance(item,Attachment) else Attachment.model_validate(item)
                # Restoring uses the persisted immutable records, never changed source paths.
                self.attachments[attachment.id] = attachment.model_copy(deep=True)
                for source in attachment.sources:
                    if re.fullmatch(r'F\d+',source.id):
                        self._source_counter = max(self._source_counter,int(source.id[1:]))

    def preview(self, attachment_id) -> Attachment:
        try:
            return self.attachments[attachment_id].model_copy(deep=True)
        except KeyError:
            raise AppError('permission', 'This attachment is not approved in this session.') from None

    def search(self, attachment_ids, query, limit=6) -> list[SourceRecord]:
        if not isinstance(limit,int) or isinstance(limit,bool) or not 1 <= limit <= 20:
            raise AppError('validation','Search limit must be 1-20.')
        if not isinstance(query,str) or len(query)>10000:
            raise AppError('validation','Search query must be text up to 10,000 characters.')
        tokens = set(re.findall(r'\w+',query.casefold()))
        candidates=[]
        for attachment_id in attachment_ids:
            for source in self.preview(attachment_id).sources:
                words = re.findall(r'\w+',source.text.casefold())
                score=sum(min(words.count(token),5) for token in tokens)/max(1,math.sqrt(len(words)))
                if score or not tokens:
                    candidates.append((score,source))
        candidates.sort(key=lambda item:(-item[0],item[1].id))
        # Retrieval is always labelled a subset, even if the small selection happens to fit.
        return [source.model_copy(update={'partial':True}) for _,source in candidates[:limit]]

    def read_range(self, attachment_id, locator) -> list[SourceRecord]:
        attachment=self.preview(attachment_id)
        if not isinstance(locator,str) or not locator or len(locator)>2000:
            raise AppError('validation','Provide an existing source ID or a page/line/cell locator.')
        matches=[]
        line_match=re.fullmatch(r'lines?\s+(\d+)(?:-(\d+))?',locator,re.I)
        page_match=re.fullmatch(r'page\s+(\d+)',locator,re.I)
        if line_match:
            start,end=int(line_match[1]),int(line_match[2] or line_match[1])
            if start<1 or end<start:
                raise AppError('validation','Line range must be positive and increasing.')
            for source in attachment.sources:
                if source.locator.startswith('lines ') and any((m:=re.match(r'(\d+): ',line)) and start<=int(m[1])<=end for line in source.text.splitlines()):
                    matches.append(source)  # Whole immutable chunks; actual supplied coverage remains visible.
        elif page_match:
            matches=[s for s in attachment.sources if re.match(r'page '+re.escape(page_match[1])+r'(?:;|$)',s.locator)]
        else:
            matches=[s for s in attachment.sources if s.id==locator or locator.casefold() in s.locator.casefold()]
        if not matches:
            raise AppError('validation','Range is not present in this immutable attachment selection. Inspect its available source locations.')
        if len(matches)>20:
            raise AppError('parser_limit','Range matches more than 20 excerpts; choose a narrower location.')
        return matches

    def table(self, attachment_id, operation, arguments) -> dict:
        attachment=self.preview(attachment_id)
        tables=attachment.metadata.get('tables',[])
        arguments=arguments or {}
        if operation not in {'count','filter','group','summary'}:
            raise AppError('validation','Table operation must be count, filter, group or summary.')
        if not isinstance(arguments,dict) or set(arguments)-{'sheet','column','group_by','where','limit'}:
            raise AppError('validation','Unexpected table arguments.')
        if not tables:
            raise AppError('unsupported','This attachment has no supported tabular selection.')
        if arguments.get('sheet'):
            tables=[table for table in tables if table['name']==arguments['sheet']]
            if not tables:
                raise AppError('validation','Selected table/sheet does not exist.')
        if len(tables)!=1:
            raise AppError('validation','Choose a table/sheet by name: '+', '.join(t['name'] for t in tables))
        table=tables[0]; headers=table['headers']; indexed=list(zip(table['row_references'],table['rows']))
        def column(name):
            if name not in headers:
                raise AppError('validation','Column must match an existing header: '+', '.join(headers))
            return headers.index(name)
        if arguments.get('where') is not None:
            where=arguments['where']
            if not isinstance(where,dict) or set(where)-{'column','op','value'} or not {'column','op'}<=set(where):
                raise AppError('validation','Filter requires column, op and optional value.')
            position=column(where['column']); op=where['op']; value=where.get('value')
            if op not in {'eq','ne','gt','gte','lt','lte','contains','missing'}:
                raise AppError('validation','Unsupported filter comparison.')
            def keep(row):
                cell=row[position] if position<len(row) else None
                if op=='missing': return cell in (None,'')
                if op=='eq': return cell==value or str(cell)==str(value)
                if op=='ne': return not (cell==value or str(cell)==str(value))
                if op=='contains': return str(value).casefold() in str(cell).casefold()
                a,b=_number(cell),_number(value)
                if a is None or b is None: return False
                return {'gt':a>b,'gte':a>=b,'lt':a<b,'lte':a<=b}[op]
            indexed=[(ref,row) for ref,row in indexed if keep(row)]
        response={'attachment_id':attachment_id,'table':table['name'],'operation':operation,'coverage':table.get('coverage','all selected rows'),'selected_rows':len(table['rows']),'matched_rows':len(indexed),'source_ids':[s.id for s in attachment.sources],'warnings':attachment.warnings}
        if operation=='count':
            response['count']=len(indexed)
        elif operation=='filter':
            limit=arguments.get('limit',100)
            if not isinstance(limit,int) or isinstance(limit,bool) or not 1<=limit<=1000:
                raise AppError('validation','Filter display limit must be 1-1000.')
            response.update(headers=headers,rows=[row for _,row in indexed[:limit]],row_references=[ref for ref,_ in indexed[:limit]],truncated=len(indexed)>limit)
        elif operation=='summary':
            position=column(arguments.get('column'))
            response.update(column=headers[position],statistics=_summary([row[position] if position<len(row) else None for _,row in indexed]))
        else:
            position=column(arguments.get('group_by')); groups={}
            numeric_position=column(arguments['column']) if arguments.get('column') else None
            for ref,row in indexed:
                key=row[position] if position<len(row) else None
                normalized=json.dumps(key,sort_keys=True)
                if normalized not in groups: groups[normalized]={'value':key,'count':0,'_numbers':[]}
                groups[normalized]['count']+=1
                if numeric_position is not None: groups[normalized]['_numbers'].append(row[numeric_position] if numeric_position<len(row) else None)
            if len(groups)>1000:
                raise AppError('parser_limit','Grouping exceeds 1,000 groups. Filter rows first.')
            for group in groups.values():
                numbers=group.pop('_numbers')
                if numeric_position is not None: group['statistics']=_summary(numbers)
            response.update(group_by=headers[position],groups=list(groups.values()))
            if numeric_position is not None:
                response['column'] = headers[numeric_position]
        return response

    def write_document(self, destination: str, content: str, format: str = 'md', overwrite: bool = False) -> dict:
        """Save literal UTF-8 document content, never execute or interpret it."""
        if not isinstance(overwrite, bool):
            raise AppError('validation', 'Overwrite must be an explicit boolean.')
        if format not in {'md', 'txt'}:
            raise AppError('validation', 'Document format must be md or txt.')
        if not isinstance(destination, str) or not destination or len(destination) > 4096:
            raise AppError('validation', 'Provide a document filename within the selected workspace.')
        if any(ord(c) < 32 or ord(c) == 127 for c in destination) or '..' in Path(destination).parts:
            raise AppError('permission', 'Document paths cannot contain control characters or parent traversal.')
        if not isinstance(content, str):
            raise AppError('validation', 'Document content must be text.')
        try:
            data = content.encode('utf-8')
        except UnicodeError:
            raise AppError('validation', 'Document content must be valid UTF-8 text.') from None
        if len(data) > MAX_DOCUMENT_BYTES:
            raise AppError('parser_limit', 'Document exceeds the 256 KiB output limit. Split it into smaller documents.')
        relative = _relative(self.workspace, destination)
        path = self.workspace / relative
        if path.suffix.lower() != '.' + format:
            raise AppError('validation', 'Document filename extension must match its md or txt format.')
        if not path.parent.is_dir():
            raise AppError('validation', 'Document destination must have an existing parent directory.')
        parent_fd = _open_at(self.workspace, relative.parent, directory=True)
        try:
            with self._lock:
                if any(a.path == str(path) for a in self.attachments.values()):
                    raise AppError('permission', 'Document cannot overwrite an attached source. Choose a new destination.')
                try:
                    current = os.stat(path.name, dir_fd=parent_fd, follow_symlinks=False)
                    if not stat.S_ISREG(current.st_mode):
                        raise AppError('permission', 'Document destination is not a regular file.')
                    if not overwrite:
                        raise AppError('validation', 'Destination exists. Choose a new document filename.')
                except FileNotFoundError:
                    pass
                _publish_bytes(parent_fd, path.name, data, overwrite)
        except OSError as exc:
            raise AppError('storage', 'Could not save document. Check directory permissions and available space.') from exc
        finally:
            os.close(parent_fd)
        return OutputArtifact(path=str(path), format=format, size_bytes=len(data),
                              sha256=hashlib.sha256(data).hexdigest()).model_dump()

    def export(self, conversation, destination, format='md', overwrite=False) -> str:
        if format not in {'md','txt','json','csv'}:
            raise AppError('validation','Export format must be md, txt, json or csv.')
        relative=_relative(self.workspace,destination)
        path=self.workspace/relative
        if not path.name or not path.parent.is_dir():
            raise AppError('validation','Export destination must have an existing parent directory.')
        # Hold parent descriptor through publication; no symlink parent traversal.
        parent_fd=_open_at(self.workspace,relative.parent,directory=True)
        try:
            try:
                current=os.stat(path.name,dir_fd=parent_fd,follow_symlinks=False)
                if not stat.S_ISREG(current.st_mode):
                    raise AppError('permission','Export destination is not a regular file.')
                if not overwrite:
                    raise AppError('validation','Destination exists. Confirm overwrite explicitly or choose another name.')
            except FileNotFoundError:
                pass
            if any(a.path==str(path) for a in self.attachments.values()):
                raise AppError('permission','Export cannot overwrite an attached source. Choose a new destination.')
            value=conversation.model_dump(mode='json') if hasattr(conversation,'model_dump') else copy.deepcopy(conversation)
            if not isinstance(value,dict):
                raise AppError('validation','Conversation export requires an object.')
            if format=='json':
                text=json.dumps(value,ensure_ascii=False,indent=2)
            elif format=='csv':
                out=io.StringIO(); writer=csv.writer(out)
                if value.get('tool_results'):
                    # A CSV has one rectangular schema: export the most recent deterministic
                    # table; JSON/Markdown/text preserve every saved result and the chat.
                    entry=value['tool_results'][-1]
                    headers, rows = _result_rows(entry['result'])
                    provenance={sid:{key:source.get(key) for key in ('title','locator','fingerprint','partial')} for sid,source in value.get('sources',{}).items() if sid in entry['result'].get('source_ids',[])}
                    attachment=self.attachments.get(entry['result'].get('attachment_id',''))
                    metadata={'created_at':entry.get('created_at',entry.get('timestamp')),'arguments':entry.get('arguments',{}),
                              'selected_rows':entry['result'].get('selected_rows'),'matched_rows':entry['result'].get('matched_rows'),
                              'truncated':entry['result'].get('truncated',False),'warnings':entry['result'].get('warnings',[]),
                              'sources':provenance,'csv_scope':'Most recent deterministic table result; JSON/Markdown/text include every saved result.'}
                    if attachment:
                        metadata.update(fingerprint=attachment.fingerprint,selection=attachment.selection)
                    names=[]
                    for name in ('record_type','source_reference','attachment_id','table','operation','coverage','source_ids','provenance_json'):
                        candidate='__hpc_'+name
                        while candidate in headers or candidate in names:
                            candidate+='_'  # Preserve actual user headers, including unusual names.
                        names.append(candidate)
                    writer.writerow([_csv_safe(h) for h in [*headers,*names]])
                    refs=entry['result'].get('row_references',[])
                    for index,row in enumerate(rows or [[None]*len(headers)]):
                        record_type='result' if rows else 'metadata_only_no_matching_rows'
                        result=entry['result']
                        extra=[record_type,refs[index] if index<len(refs) else '',result.get('attachment_id',''),result.get('table',''),
                               result.get('operation',entry.get('operation','')),result.get('coverage',''),' '.join(result.get('source_ids',[])),
                               json.dumps(metadata,ensure_ascii=False) if index==0 else '']
                        writer.writerow([_csv_safe(cell) for cell in [*row,*extra]])
                else:
                    writer.writerow(['turn_id','user','assistant','source_ids','source_provenance'])
                    for turn in value.get('turns',[]):
                        sources=turn.get('sources',[])
                        writer.writerow([_csv_safe(turn.get('id','')),_csv_safe(turn.get('request',{}).get('text','')),_csv_safe(turn.get('answer','')),
                                         ' '.join(s['id'] for s in sources),_csv_safe(json.dumps(sources,ensure_ascii=False))])
                text=out.getvalue()
            else:
                lines=[clean_text(value.get('title','Conversation')),'']
                for turn in value.get('turns',[]):
                    lines.extend(['User: '+turn.get('request',{}).get('text',''),'','Assistant: '+turn.get('answer',''),''])
                if value.get('tool_results'):
                    lines.extend(['Deterministic table results (computed locally from immutable selected data):',''])
                    for index,entry in enumerate(value['tool_results'],1):
                        lines.extend([f"Result {index}: {entry.get('operation',entry.get('result',{}).get('operation','table'))}",
                                      json.dumps(entry,ensure_ascii=False,indent=2),''])
                sources=dict(value.get('sources',{}))
                for turn in value.get('turns',[]):
                    sources.update({s['id']:s for s in turn.get('sources',[])})
                lines.extend(['Sources (immutable evidence snapshots; citation existence does not establish factual entailment):',''])
                for sid,source in sources.items():
                    lines.extend([f"[{sid}] {source.get('title','')} — {source.get('locator','')}",
                                  f"URL: {source.get('url') or ''}; retrieved: {source.get('retrieved_at','')}; SHA256: {source.get('fingerprint','')}; partial: {source.get('partial',False)}",source.get('text',''),''])
                text='\n'.join(lines)
            if len(text.encode())>32*1024*1024:
                raise AppError('parser_limit','Export exceeds 32 MiB; export a smaller conversation.')
            _publish_bytes(parent_fd, path.name, text.encode('utf-8'), overwrite)
            return str(path)
        finally:
            os.close(parent_fd)


def _result_rows(result: dict) -> tuple[list[str],list[list]]:
    """Normalize the four deterministic operations without asking a model to format data."""
    operation=result.get('operation')
    if operation=='filter':
        return result.get('headers',[]),result.get('rows',[])
    if operation=='count':
        return ['count'],[[result.get('count')]]
    if operation=='summary':
        statistics=result.get('statistics',{})
        headers=['column',*statistics]
        return headers,[[result.get('column'),*statistics.values()]]
    if operation=='group':
        groups=result.get('groups',[])
        statistics=list(dict.fromkeys(key for group in groups for key in group.get('statistics',{})))
        headers=['group_by','group_value','count']
        if result.get('column'):
            headers.append('numeric_column')
        headers.extend(statistics)
        rows=[]
        for group in groups:
            row=[result.get('group_by'),group.get('value'),group.get('count')]
            if result.get('column'):
                row.append(result['column'])
            row.extend(group.get('statistics',{}).get(key) for key in statistics)
            rows.append(row)
        return headers,rows
    raise AppError('validation','Saved deterministic result has an unknown operation.')


def _number(value):
    if value is None or isinstance(value,bool) or (isinstance(value,str) and not value.strip()): return None
    try:
        number=float(value)
        return number if math.isfinite(number) else None
    except (ValueError,TypeError): return None


def _summary(values):
    numbers=[number for value in values if (number:=_number(value)) is not None]
    return {'numeric_count':len(numbers),'missing_count':sum(v is None or v=='' for v in values),
            'nonnumeric_count':sum(v is not None and v!='' and _number(v) is None for v in values),
            'sum':sum(numbers) if numbers else None,'mean':sum(numbers)/len(numbers) if numbers else None,
            'min':min(numbers) if numbers else None,'max':max(numbers) if numbers else None,
            'note':'Original columns and units are preserved; numeric parsing accepts finite plain numbers only.'}


def _csv_safe(value):
    if value is None:
        return ''
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        return value
    text=clean_text(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text


def _worker() -> None:
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(1024*1024*1024,1024*1024*1024))
        resource.setrlimit(resource.RLIMIT_CPU,(10,10))
        resource.setrlimit(resource.RLIMIT_FSIZE,(32*1024*1024,32*1024*1024))
        payload=json.loads(sys.stdin.buffer.read(24*1024*1024))
        data=base64.b64decode(payload['data'],validate=True)
        parsed=_parse(data,payload['extension'],payload['selection'],payload.get('vision',False))
        serialized=json.dumps(parsed,ensure_ascii=False,allow_nan=False)
        if len(serialized.encode('utf-8')) > 32*1024*1024:
            raise AppError('parser_limit','Extracted document exceeds the 32 MiB structured-output limit; select a smaller range.')
        print(serialized)
    except AppError as exc:
        print(json.dumps({'error':exc.message,'code':exc.code}))
    except (MemoryError,RecursionError):
        print(json.dumps({'error':'Parser memory or nesting limit exceeded.','code':'parser_limit'}))
    except Exception:
        print(json.dumps({'error':'Malformed document or unsupported content; no content was attached.','code':'validation'}))


if __name__=='__main__' and '--parse' in sys.argv:
    _worker()
